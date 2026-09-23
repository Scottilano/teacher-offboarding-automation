"""Read-only extraction. No macros, formulas, external links or file writes."""
import csv
import io
import json
import sys
import zipfile

MAX_BYTES = 50 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 100


def extract(raw, extension, requested_sheet, max_bytes=MAX_BYTES):
    if not raw or len(raw) > max_bytes:
        raise ValueError(f"The file is empty or exceeds {max_bytes // (1024 * 1024)} MB.")
    if extension == ".csv":
        encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        text = raw.decode(encoding, errors="strict")
        csv.field_size_limit(32768)
        rows = []
        for row in csv.reader(io.StringIO(text, newline=""), strict=True):
            if len(rows) >= MAX_ROWS or len(row) > MAX_COLUMNS:
                raise ValueError("Rosters are limited to 10000 rows and 100 columns. Reduce the file size.")
            if any("\x00" in cell for cell in row):
                raise ValueError("CSV contains invalid characters. Export a new UTF-8 CSV.")
            rows.append(row)
        return {"sheetNames": [], "sheetName": "CSV", "rows": rows}
    if extension != ".xlsx":
        raise ValueError("Only .xlsx and .csv are supported. Save legacy .xls files as .xlsx.")

    # Bound decompression and reject XML declarations capable of entity expansion.
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(i.file_size for i in entries) > 4 * max_bytes:
            raise ValueError("Excel decompressed size exceeds the limit. Export only the roster.")
        if len({i.filename for i in entries}) != len(entries):
            raise ValueError("Excel contains duplicate ZIP entries.")
        for entry in entries:
            if entry.file_size > 4 * max_bytes or entry.flag_bits & 1:
                raise ValueError("Oversized or encrypted Excel files are not supported.")
            if entry.filename.endswith((".xml", ".rels")):
                xml = archive.read(entry).replace(b"\x00", b"").upper()
                if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                    raise ValueError("Excel XML contains prohibited entity declarations.")
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:
        names = workbook.sheetnames
        if requested_sheet is None and len(names) != 1:
            return {"sheetNames": names, "needsSheet": True}
        name = requested_sheet if requested_sheet is not None else names[0]
        if name not in names:
            raise ValueError("The specified worksheet does not exist.")
        sheet = workbook[name]
        if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
            raise ValueError("Worksheets are limited to 10000 rows and 100 columns. Remove extra rows or columns and retry.")
        sheet.reset_dimensions()  # Do not silently trust an incorrect XML dimension.
        rows = []
        for index, cells in enumerate(sheet.iter_rows()):
            if index >= MAX_ROWS or len(cells) > MAX_COLUMNS:
                raise ValueError("The actual worksheet dimensions exceed the limit.")
            row = []
            for cell in cells:
                value = cell.value
                if cell.data_type in ("f", "e"):
                    row.append({"invalid": "FORMULA_OR_ERROR"})
                elif value is None:
                    row.append("")
                elif isinstance(value, (str, int, float, bool)):
                    row.append(value)
                else:
                    row.append(str(value))
            rows.append(row)
        return {"sheetNames": names, "sheetName": name, "rows": rows}
    finally:
        workbook.close()


if __name__ == "__main__":
    try:
        max_bytes = int(sys.argv[3]) if len(sys.argv) > 3 else MAX_BYTES
        if not 1024 * 1024 <= max_bytes <= 100 * 1024 * 1024:
            raise ValueError("Invalid file-size limit configuration.")
        raw = sys.stdin.buffer.read(max_bytes + 1)
        sheet = (sys.argv[2] or None) if len(sys.argv) > 2 else None
        print(json.dumps(extract(raw, sys.argv[1], sheet, max_bytes), ensure_ascii=False))
    except Exception as error:
        # No source cell contents or credentials in parser errors.
        message = str(error) if isinstance(error, ValueError) else "Unable to read the file. Export a new UTF-8 CSV or an unencrypted .xlsx file."
        print(json.dumps({"error": message}, ensure_ascii=False))
        sys.exit(1)
